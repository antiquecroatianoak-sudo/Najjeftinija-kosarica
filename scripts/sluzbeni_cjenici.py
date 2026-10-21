"""
Službeni cjenici lanaca (NN 101/2026: t. II XML ili CSV svaki dan do 8:00, t. V trgovci s internetskom stranicom,
t. VII automatsko preuzimanje). Za svaki lanac s javne stranice preuzima današnje datoteke po poslovnici, čita retke
(naziv, šifra, količina, jedinica, MPC, posebna cijena, cijena za jedinicu mjere, barkod) i za svaki redak čuva trag:
URL datoteke, redni broj retka/zapisa i datum.

Ne zaobilazi zaštitu: poštuje robots.txt; 403, JS izazov, prijava ili ključ (Google Drive) bilježe se kao neuspjeh
s razlogom. Lanac koji danas ne uspije ostaje iz arhive cijene.dev i označen je kao neprovjeren prema službenom izvoru.

Samostalno: python scripts/sluzbeni_cjenici.py [GGGG-MM-DD] [lanac ...] ispisuje pokrivenost (bez zapisivanja).
"""
import csv
import html
import io
import json
import re
import sys
import threading
import time
import unicodedata
import urllib.parse
import urllib.request
import urllib.robotparser
import xml.etree.ElementTree as ET
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from zoneinfo import ZoneInfo

AGENT = "NajjeftinijaKosarica/1.0 (+https://github.com/antiquecroatianoak-sudo/Najjeftinija-kosarica)"
ISTOVREMENO = 6           # preuzimanja odjednom, ukupno
ISTOVREMENO_PO_POSLUZITELJU = 3
NAJVISE_BAJTOVA = 200_000_000

# Stupci službenih cjenika (normalizirani naziv zaglavlja ili XML oznake) -> polje.
STUPCI = {
    "naziv": ["naziv proizvoda", "naziv", "naziv artikla", "naziv art", "proizvod", "artikl"],
    "sifra": ["sifra proizvoda", "sifra", "sif art", "sifra artikla", "kod proizvoda", "sifra art"],
    "marka": ["marka proizvoda", "marka", "brand"],
    "kolicina": ["neto kolicina", "net kol", "nettokolicina", "neto kol", "gramaza", "neto kolicina proizvoda", "kolicina",
                 "neto kolicina kg"],
    # (stvarna zaglavlja iz pokretanja 21.10.2026: Spar "(EUR)", Plodine "Cijena po JM", Tommy, Stridon, Eurospin)
    "jedinica": ["jedinica mjere", "jmj", "mjera", "jmjere", "j mjere", "jedinica", "jm", "jed mj 1 kom l kg"],
    "cijena": ["maloprodajna cijena", "mpc", "mpcijena", "maloprod cijena", "cijena"],
    "jedinicna": ["cijena za jedinicu mjere", "c jmj", "mpcijenamjera", "mpcjmj", "cijena jedinice mjere",
                  "jedinicna cijena", "cijena po jedinici mjere", "mpc jmj", "cijena po jm", "cijena po jmj", "cijena jed mj"],
    "posebna": ["mpc za vrijeme posebnog oblika prodaje", "mpc pop", "akcijska cijena", "posebna cijena",
                "mpc posebnog oblika prodaje", "mpc posebni oblik prodaje", "mpc posebna prodaja", "mpc poseb oblik prod",
                "maloprodajna cijena akcija", "akc cijena a akcija"],
    "akcija": ["naziv posebnog oblika prodaje", "poseban oblik prodaje"],
    "barkod": ["barkod", "ean kod", "barcode", "ean", "bar kod", "barkod artikla"],
    "dostupnost": ["dostupnost", "raspolozivost", "dostupno nedostupno", "dostupno", "dostupan nedostupan"],
}
INDEKS_STUPACA = {n: k for k, nazivi in STUPCI.items() for n in nazivi}


def normaliziraj(tekst: str) -> str:
    tekst = (tekst or "").lower().replace("đ", "d")
    tekst = "".join(z for z in unicodedata.normalize("NFD", tekst) if not unicodedata.combining(z))
    return re.sub(r"[^a-z0-9]+", " ", tekst).strip()


def kljuc_stupca(ime: str) -> str | None:
    # XML oznake u CamelCase (Studenac, Ribola: "NazivProizvoda", "MaloprodajnaCijena") razdvajaju se u riječi.
    # Samo prava CamelCase oznaka (jedna riječ, mješovita slova); "NETO KOLIčINA" (Dukat) ostaje kakva jest.
    if re.fullmatch(r"[A-Za-zČĆŠŽĐčćšžđ]+", ime) and not ime.isupper() and not ime.islower():
        ime = re.sub(r"(?<=[a-zčćšžđ])(?=[A-ZČĆŠŽĐ])|(?<=[A-ZČĆŠŽĐ])(?=[A-ZČĆŠŽĐ][a-zčćšžđ])", " ", ime)
    n = normaliziraj(ime)
    # Valuta u zaglavlju ("MPC (EUR)", "MALOPROD.CIJENA(EUR)") ne mijenja značenje stupca.
    n = re.sub(r"\s+eur$", "", n)
    if n in INDEKS_STUPACA:
        return INDEKS_STUPACA[n]
    # "SIDRENA CIJENA NA 2.5.2025", "NAJNIŽA CIJENA U POSLJEDNIH 30 DANA" i sl. nisu cijena ponude.
    return None


# ---------- Preuzimanje ----------
_roboti: dict = {}
_brave = threading.Lock()
_posluzitelji: dict[str, threading.Semaphore] = {}


def siguran_url(url: str) -> str:
    """Razmaci i dijakritike u poveznicama (Lidl, Roto, Stridon, Boso) kodiraju se; postojeći %xx ostaju."""
    p = urllib.parse.urlsplit(html.unescape(url))
    put = urllib.parse.quote(p.path, safe="/%:@&=+$,;!~*'()")
    upit = urllib.parse.quote(p.query, safe="=&%:/,;+@!$'()*~")
    return urllib.parse.urlunsplit((p.scheme, p.netloc.encode("idna").decode() if p.netloc else "", put, upit, ""))


def dopusteno(url: str) -> bool:
    p = urllib.parse.urlsplit(url)
    korijen = f"{p.scheme}://{p.netloc}"
    with _brave:
        rp = _roboti.get(korijen)
    if rp is None:
        rp = urllib.robotparser.RobotFileParser()
        try:
            rp.parse(_dohvati(korijen + "/robots.txt", 1_000_000).decode("utf-8", "replace").splitlines())
        except Exception:  # noqa: BLE001 - bez robots.txt vrijedi dopušteno
            rp.parse([])
        with _brave:
            _roboti[korijen] = rp
    return rp.can_fetch(AGENT, url)


def _dohvati(url: str, najvise: int = NAJVISE_BAJTOVA) -> bytes:
    zahtjev = urllib.request.Request(url, headers={"User-Agent": AGENT, "Accept": "*/*"})
    with urllib.request.urlopen(zahtjev, timeout=90) as r:
        sadrzaj = r.read(najvise + 1)
    if len(sadrzaj) > najvise:
        raise ValueError("datoteka je prevelika")
    return sadrzaj


def dohvati(url: str) -> bytes:
    url = siguran_url(url)
    if not dopusteno(url):
        raise PermissionError("robots.txt ne dopušta")
    host = urllib.parse.urlsplit(url).netloc
    with _brave:
        sem = _posluzitelji.setdefault(host, threading.Semaphore(ISTOVREMENO_PO_POSLUZITELJU))
    with sem:
        for pokusaj in range(3):
            try:
                sadrzaj = _dohvati(url)
                break
            except (PermissionError, ValueError):
                raise
            except Exception as e:  # noqa: BLE001
                msg = str(e)
                if "403" in msg or pokusaj == 2 or ("404" in msg and pokusaj == 1):
                    raise
                time.sleep(3 * (pokusaj + 1))
    pocetak = sadrzaj[:600].lower()
    if b"just a moment" in pocetak or b"cf-challenge" in pocetak or b"captcha" in pocetak:
        raise PermissionError("JS izazov / zaštita od botova (ne zaobilazi se)")
    return sadrzaj


def dekodiraj(sadrzaj: bytes) -> str:
    for kod in ("utf-8-sig", "cp1250"):
        try:
            return sadrzaj.decode(kod)
        except UnicodeDecodeError:
            continue
    return sadrzaj.decode("latin-1")


# ---------- Čitanje datoteka ----------
def citaj_csv(tekst: str):
    """(zaglavlje, nepoznati stupci, zapisi s _redak = redak u datoteci, zaglavlje je 1. redak)."""
    uzorak = tekst[:8192]
    try:
        dijalekt = csv.Sniffer().sniff(uzorak, delimiters=";,\t|")
    except csv.Error:
        dijalekt = csv.excel
        dijalekt.delimiter = ";" if uzorak.count(";") > uzorak.count(",") else ","
    redovi = csv.reader(io.StringIO(tekst), dijalekt)
    zaglavlje = next(redovi, [])
    kljucevi = [kljuc_stupca(z) for z in zaglavlje]
    nepoznati = [z.strip() for z, k in zip(zaglavlje, kljucevi) if not k and z.strip()]
    zapisi = []
    if "naziv" in kljucevi and "cijena" in kljucevi:
        for red in redovi:
            zapis = {k: v.strip() for k, v in zip(kljucevi, red) if k and k not in ("",)}
            if zapis.get("naziv"):
                zapis["_redak"] = redovi.line_num
                zapisi.append(zapis)
    return zaglavlje, nepoznati, zapisi


def citaj_xml(tekst: str):
    """Zapis = element čija su djeca listovi s prepoznatim nazivom i cijenom; _redak = redni broj zapisa."""
    tekst = re.sub(r"&(?![A-Za-z][A-Za-z0-9]*;|#\d+;|#x[0-9A-Fa-f]+;)", "&amp;", tekst)
    korijen = ET.fromstring(tekst.encode("utf-8") if not tekst.lstrip().startswith("<?xml") else
                            re.sub(r"^<\?xml[^>]*\?>", "", tekst.lstrip()).encode("utf-8"))
    zapisi, nepoznati, oznake = [], set(), []
    for el in korijen.iter():
        djeca = list(el)
        if not djeca or any(list(d) for d in djeca):
            continue
        zapis = {}
        for d in djeca:
            ime = d.tag.split("}")[-1]
            k = kljuc_stupca(ime.replace("_", " "))
            if k and (d.text or "").strip():
                zapis[k] = d.text.strip()
            elif not k:
                nepoznati.add(ime)
        if zapis.get("naziv") and "cijena" in zapis or (zapis.get("naziv") and zapis.get("posebna")):
            zapis["_redak"] = len(zapisi) + 1
            zapisi.append(zapis)
            if not oznake:
                oznake = [d.tag.split("}")[-1] for d in djeca]
    return oznake, sorted(nepoznati), zapisi


def citaj_datoteku(sadrzaj: bytes, ime: str):
    if sadrzaj[:2] == b"PK" and not ime.lower().endswith(".zip"):
        return citaj_xlsx(sadrzaj)
    tekst = dekodiraj(sadrzaj)
    if tekst.lstrip().startswith("<"):
        oznake, nepoznati, zapisi = citaj_xml(tekst)
        if not zapisi:
            oznake = [f"početak XML-a: {re.sub(r'\s+', ' ', tekst[:260])}"]
        return oznake, nepoznati, zapisi
    return citaj_csv(tekst)


def citaj_xlsx(sadrzaj: bytes):
    """Prvi list XLSX-a: prvi redak je zaglavlje (Djelo Vodice, dm). Treba paket openpyxl."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(sadrzaj), read_only=True, data_only=True)
    redovi = wb.worksheets[0].iter_rows(values_only=True)
    zaglavlje = [str(x or "").strip() for x in next(redovi, [])]
    # Neki listovi imaju naslov iznad zaglavlja: zaglavlje je prvi redak s prepoznatim nazivom i cijenom.
    preskoceno = 1
    while not {"naziv", "cijena"} <= {kljuc_stupca(z) for z in zaglavlje} and preskoceno < 60:
        zaglavlje = [str(x or "").strip() for x in next(redovi, [])]
        preskoceno += 1
    kljucevi = [kljuc_stupca(z) for z in zaglavlje]
    nepoznati = [z for z, k in zip(zaglavlje, kljucevi) if not k and z]
    zapisi = []
    if not ("naziv" in kljucevi and "cijena" in kljucevi):
        prvi = [[str(x) for x in r[:12] if x is not None] for r in wb.worksheets[0].iter_rows(max_row=8, values_only=True)]
        return [f"XLSX listovi {wb.sheetnames[:3]}, prvi retci {prvi}"], nepoznati, []
    if "naziv" in kljucevi and "cijena" in kljucevi:
        for n, red in enumerate(redovi, start=preskoceno + 1):
            zapis = {k: ("" if v is None else str(v)).strip() for k, v in zip(kljucevi, red) if k}
            if zapis.get("naziv"):
                zapis["_redak"] = n
                zapisi.append(zapis)
    return zaglavlje, nepoznati, zapisi


# ---------- Datumi i poslovnice iz naziva datoteka ----------
def oblici_datuma(d: date) -> list[str]:
    return [f"{d:%d.%m.%Y}", f"{d:%d_%m_%Y}", f"{d:%d-%m-%Y}", f"{d:%d%m%Y}", f"{d:%Y%m%d}", f"{d:%Y-%m-%d}",
            f"{d.day}.{d.month}.{d.year}", f"{d:%d.%m.%y}"]


def ima_datum(ime: str, d: date) -> bool:
    ime = urllib.parse.unquote(ime)
    return any(o in ime for o in oblici_datuma(d))


def kljuc_poslovnice(ime: str, d: date | None = None) -> str:
    """
    Naziv datoteke (ili naslov) do današnjeg datuma, bez nastavka: ista poslovnica u istom danu i kad lanac objavi više
    datoteka (dopune s vremenom u nazivu). Bez datuma u nazivu cijeli je naziv poslovnica.
    """
    ime = urllib.parse.unquote(ime.rsplit("/", 1)[-1].split("#")[-1])
    ime = re.sub(r"\.(csv|xml|zip|xlsx?)$", "", ime, flags=re.I)
    if d is not None:
        mjesta = [ime.find(o) for o in oblici_datuma(d) if o in ime]
        if mjesta:
            ime = ime[:min(mjesta)]
    return ime.strip(" ,_.#-") or ime


def opis_poslovnice(kljuc: str) -> tuple[str, str]:
    """(adresa, grad) najbolje što se iz naziva da: razdvaja zarezima, podvlakama ili #."""
    dijelovi = [x.strip() for x in re.split(r"[,_#]+", kljuc) if x.strip()]
    tekst = " ".join(dijelovi)
    return tekst[:120], ""


def poveznice(tekst: str, osnova: str, uzorak: str = r"\.(csv|xml|zip|xlsx?)(?:$|[?#&])") -> list[str]:
    rez = []
    for m in re.finditer(r"""(?:href|value|src)\s*=\s*["']([^"']+)["']""", tekst, re.I):
        href = html.unescape(m.group(1))
        if re.search(uzorak, href, re.I):
            rez.append(urllib.parse.urljoin(osnova, href))
    return list(dict.fromkeys(rez))


# ---------- Lanci: popis današnjih datoteka [(url, ključ poslovnice)] ----------
def s_indeksa(url: str, d: date, uzorak=r"\.(csv|xml)(?:$|[?#&])"):
    tekst = dekodiraj(dohvati(url))
    return [(u, kljuc_poslovnice(u, d)) for u in poveznice(tekst, url, uzorak) if ima_datum(u, d)]


def konzum(d):
    datoteke, stranica = [], 1
    while stranica < 40:
        tekst = dekodiraj(dohvati(f"https://www.konzum.hr/cjenici?date={d:%Y-%m-%d}&page={stranica}"))
        nove = [u for u in poveznice(tekst, "https://www.konzum.hr/", r"/cjenici/download\?title=") if u not in {x for x, _ in datoteke}]
        if not nove:
            break
        for u in nove:
            naslov = urllib.parse.parse_qs(urllib.parse.urlsplit(u).query).get("title", [""])[0]
            if ima_datum(naslov, d):
                datoteke.append((u, kljuc_poslovnice(naslov, d)))
        stranica += 1
    return datoteke


def spar(d):
    indeks = json.loads(dekodiraj(dohvati(f"https://www.spar.hr/datoteke_cjenici/Cjenik{d:%Y%m%d}.json")))
    return [(f["URL"], kljuc_poslovnice(f.get("name") or f["URL"], d)) for f in indeks.get("files", []) if f.get("URL")]


def tommy(d):
    osnova = "https://spiza.tommy.hr/api/v2"
    indeks = json.loads(dekodiraj(dohvati(f"{osnova}/shop/store-prices-tables?date={d:%Y-%m-%d}&page=1&itemsPerPage=200&channelCode=general")))
    rez = []
    for clan in indeks.get("hydra:member", []):
        put = clan.get("@id", "")
        if put.startswith("/api/v2"):
            put = put[len("/api/v2"):]
        rez.append((osnova + put, kljuc_poslovnice(clan.get("fileName") or put, d)))
    return rez


def kaufland(d):
    tekst = dekodiraj(dohvati("https://www.kaufland.hr/akcije-novosti/popis-mpc.html"))
    rez = []
    for m in re.finditer(r"""data-props\s*=\s*(["'])(.*?)\1""", tekst, re.S):
        try:
            props = json.loads(html.unescape(m.group(2)))
        except json.JSONDecodeError:
            continue
        put = (props.get("settings") or {}).get("dataUrlAssets")
        if not put:
            continue
        json_url = urllib.parse.urljoin("https://www.kaufland.hr/", put)
        stavke = json.loads(dekodiraj(dohvati(json_url)))
        if stavke:
            print(f"    kaufland: popis {json_url}, prva stavka {json.dumps(stavke[0], ensure_ascii=False)[:300]}", flush=True)
        for stavka in stavke:
            vrijednosti = [str(v) for v in stavka.values()] if isinstance(stavka, dict) else [str(stavka)]
            u = next((v for v in vrijednosti if re.search(r"^(https?:|/).*\.csv", v, re.I)), None) or \
                next((v for v in vrijednosti if re.search(r"\.csv", v, re.I)), None)
            naslov = " ".join(vrijednosti)
            if u and ima_datum(naslov, d):
                # Samo naziv datoteke: nalazi se uz popis (isti direktorij kao JSON).
                rez.append((urllib.parse.urljoin(json_url, u), kljuc_poslovnice(u, d)))
    if not rez:
        raise ValueError("popis datoteka (data-props → dataUrlAssets) nije pronađen ili nema današnjih")
    return rez


def ktc(d):
    tekst = dekodiraj(dohvati("https://www.ktc.hr/cjenici"))
    rez = []
    for poslovnica in poveznice(tekst, "https://www.ktc.hr/", r"cjenici\?poslovnica="):
        for u in poveznice(dekodiraj(dohvati(poslovnica)), "https://www.ktc.hr/", r"\.csv"):
            if ima_datum(u, d):
                rez.append((u, kljuc_poslovnice(u, d)))
    return rez


def poveznice_s_tekstom(tekst: str, osnova: str, uzorak: str) -> list[tuple[str, str]]:
    """[(url, tekst oko poveznice)]: datum objave ponekad piše samo u naslovu na stranici (Žabac)."""
    rez = []
    for m in re.finditer(r"""<a\b[^>]*href\s*=\s*["']([^"']+)["'][^>]*>(.*?)</a>""", tekst, re.I | re.S):
        if re.search(uzorak, m.group(1), re.I):
            okolina = re.sub(r"<[^>]+>", " ", tekst[max(0, m.start() - 1500):m.end()])
            rez.append((urllib.parse.urljoin(osnova, html.unescape(m.group(1))), html.unescape(okolina)))
    return rez


def zabac(d):
    osnova = "https://zabacfoodoutlet.hr/cjenik/"
    stranice = [osnova] + poveznice(dekodiraj(dohvati(osnova)), osnova, r"\?store=")
    rez = {}
    for st in dict.fromkeys(stranice):
        for u, okolina in poveznice_s_tekstom(dekodiraj(dohvati(st)), st, r"\.(csv|xlsx)(?:$|[?#&])"):
            datumi = re.findall(r"\d{1,2}\.\s*\d{1,2}\.\s*\d{4}", okolina)
            if datumi and ima_datum(re.sub(r"\s+", "", datumi[-1]), d) and u not in rez:
                trgovina = urllib.parse.parse_qs(urllib.parse.urlsplit(st).query).get("store", ["Žabac"])[0]
                rez[u] = trgovina
    if not rez:
        uzorak = next((o for st in stranice[:1] for _, o in poveznice_s_tekstom(dekodiraj(dohvati(st)), st, r"\.(csv|xlsx)")), "")
        raise ValueError(f"na stranicama nema datoteke s datumom {d:%d.%m.%Y} u naslovu; tekst uz poveznicu: "
                         f"{re.sub(r'\s+', ' ', uzorak)[-200:]!r}")
    return list(rez.items())


def dm(d):
    indeks = json.loads(dekodiraj(dohvati("https://content.services.dmtech.com/rootpage-dm-shop-hr-hr/novo/promocije/"
                                          "nove-oznake-cijena-i-vazeci-cjenik-u-dm-u-2906632?mrclx=false")))
    rez = []
    for stavka in indeks.get("mainData", []):
        if stavka.get("type") != "CMDownload":
            continue
        podatak = stavka.get("data") or {}
        naslov, cilj = podatak.get("headline", ""), podatak.get("linkTarget", "")
        if cilj and ima_datum(naslov, d):
            rez.append((urllib.parse.urljoin("https://content.services.dmtech.com/rootpage-dm-shop-hr-hr", cilj), "dm"))
    if not rez:
        raise ValueError(f"nema cjenika s datumom {d:%d.%m.%Y} u naslovu")
    return rez[:1]


def plodine(d):
    return [(u, "ZIP") for u in poveznice(dekodiraj(dohvati("https://www.plodine.hr/info-o-cijenama")),
                                          "https://www.plodine.hr/", r"\.zip") if ima_datum(u, d)]


def eurospin(d):
    return [(u, "ZIP") for u in poveznice(dekodiraj(dohvati("https://www.eurospin.hr/cjenik/")),
                                          "https://www.eurospin.hr/", r"\.zip") if ima_datum(u, d)]


LANCI = {
    "konzum": konzum,
    "lidl": lambda d: s_indeksa("https://www.lidl.hr/c/cijene/s10073252", d),
    "plodine": plodine,
    "kaufland": kaufland,
    "spar": spar,
    "ntl": lambda d: s_indeksa(f"https://ntl.hr/cjenik/?date={d:%Y-%m-%d}", d),
    "studenac": lambda d: [(f"https://www.studenac.hr/cjenici/PROIZVODI-{d:%Y-%m-%d}.zip", "ZIP")],
    "tommy": tommy,
    "ktc": ktc,
    "eurospin": eurospin,
    "ribola": lambda d: s_indeksa(f"https://ribola.hr/ribola-cjenici/?date={d:%d.%m.%Y}", d),
    "boso": lambda d: s_indeksa("https://www.boso.hr/cjenik/", d),
    "trgocentar": lambda d: s_indeksa("https://trgocentar.com/Trgovine-cjenik/", d),
    "trgovina-krk": lambda d: s_indeksa("https://trgovina-krk.hr/objava-cjenika/", d),
    "vrutak": lambda d: s_indeksa("https://www.vrutak.hr/cjenik-svih-artikala", d),
    "lorenco": lambda d: s_indeksa("https://lorenco.hr/dnevne-cijene/", d),
    "roto": lambda d: s_indeksa("https://www.rotodinamic.hr/cjenici/", d),
    "stridon": lambda d: s_indeksa("https://stridon.hr/hr/supermarketi", d),
    "branka": lambda d: s_indeksa("https://www.branka.hr/cjenik", d),
    "dukat": lambda d: s_indeksa("https://dukat.hr/diskonti", d),
    "jadranka_trgovina": lambda d: s_indeksa("https://jadranka-trgovina.com/cjenici/", d),
    "bure": lambda d: s_indeksa("https://www.bure.hr/cjenici-arhiva", d),
    "brodokomerc": lambda d: s_indeksa("http://www.brodokomerc.hr/cijene", d),
    "zabac": zabac,
    "djelo_vodice": lambda d: s_indeksa("https://dv10.djelo-vodice.hr/", d, r"\.xlsx(?:$|[?#&])"),
    "dm": dm,
    "stanic": lambda d: (_ for _ in ()).throw(PermissionError("datoteke na Google Driveu traže API ključ; ne koristi se")),
}


def _datoteke_iz(url: str, kljuc: str, d: date):
    """[(url, ključ poslovnice, ime, sadržaj)]; ZIP se raspakira (član = poslovnica)."""
    sadrzaj = dohvati(url)
    if kljuc == "ZIP" or (sadrzaj[:2] == b"PK" and url.lower().split("?")[0].endswith(".zip")):
        zf = zipfile.ZipFile(io.BytesIO(sadrzaj))
        return [(f"{url}#{ime}", kljuc_poslovnice(ime, d), ime, zf.read(ime)) for ime in zf.namelist()
                if re.search(r"\.(csv|xml)$", ime, re.I)]
    return [(url, kljuc, url, sadrzaj)]


def preuzmi_lanac(lanac: str, d: date) -> dict:
    """{"lanac", "uspjeh", "razlog", "datoteke": [{"url", "poslovnica", "zapisi", "nepoznati", "greska"}]}."""
    rez = {"lanac": lanac, "uspjeh": False, "razlog": "", "datoteke": []}
    try:
        popis = LANCI[lanac](d)
    except Exception as e:  # noqa: BLE001
        rez["razlog"] = f"popis datoteka: {str(e)[:200]}"
        return rez
    if not popis:
        rez["razlog"] = f"nema datoteka s datumom {d:%d.%m.%Y}"
        return rez

    def jedna(par):
        url, kljuc = par
        try:
            izlaz = []
            for u, k, ime, sadrzaj in _datoteke_iz(url, kljuc, d):
                try:
                    zaglavlje, nepoznati, zapisi = citaj_datoteku(sadrzaj, ime)
                    izlaz.append({"url": u, "poslovnica": k, "zapisi": zapisi, "nepoznati": nepoznati,
                                  "greska": "" if zapisi else f"nema prepoznatih stupaca naziv/cijena: {str(zaglavlje[:20])[:420]}"})
                except Exception as e:  # noqa: BLE001
                    izlaz.append({"url": u, "poslovnica": k, "zapisi": [], "nepoznati": [], "greska": f"{str(e)[:200]} ({u[:160]})"})
            return izlaz
        except Exception as e:  # noqa: BLE001
            return [{"url": url, "poslovnica": kljuc, "zapisi": [], "nepoznati": [], "greska": f"{str(e)[:200]} ({siguran_url(url)[:160]})"}]

    with ThreadPoolExecutor(ISTOVREMENO) as ex:
        for izlaz in ex.map(jedna, popis):
            rez["datoteke"].extend(izlaz)
    # Ista poslovnica više puta u danu (dopune): vrijedi zadnja datoteka po nazivu (vrijeme je u nazivu).
    po_poslovnici = {}
    for dat in sorted(rez["datoteke"], key=lambda x: x["url"]):
        if dat["zapisi"]:
            po_poslovnici[dat["poslovnica"]] = dat
    rez["koristene"] = list(po_poslovnici.values())
    rez["uspjeh"] = bool(rez["koristene"])
    if not rez["uspjeh"]:
        rez["razlog"] = "nijedna datoteka nije pročitana: " + "; ".join({d["greska"] for d in rez["datoteke"]})[:300]
    return rez


def u_arhivu_lanca(rez: dict, d: date) -> tuple[str, str, str, list]:
    """Službeni podaci jednog lanca u obliku arhive cijene.dev (stores/products/prices.csv) + popis datoteka (izvori)."""
    trgovine, proizvodi, cijene, izvori = io.StringIO(), {}, io.StringIO(), []
    wt = csv.writer(trgovine)
    wt.writerow(["store_id", "type", "address", "city", "zipcode"])
    wc = csv.writer(cijene)
    wc.writerow(["store_id", "product_id", "price", "unit_price", "special_price", "akcija", "izvor", "redak"])
    for dat in rez["koristene"]:
        izvor = len(izvori)
        osnova, _, clan = dat["url"].partition("#")
        izvori.append([siguran_url(osnova) + (f"#{clan}" if clan else ""), d.isoformat()])
        adresa, grad = opis_poslovnice(dat["poslovnica"])
        store_id = dat["poslovnica"]
        wt.writerow([store_id, "", adresa, grad, ""])
        for z in dat["zapisi"]:
            if normaliziraj(z.get("dostupnost", "")) in ("nedostupno", "nedostupan", "ne"):
                continue
            pid = z.get("sifra") or z.get("barkod") or normaliziraj(z["naziv"])
            if pid not in proizvodi:
                proizvodi[pid] = [pid, (z.get("barkod") or "").strip(), z["naziv"], z.get("marka", ""),
                                  z.get("kolicina", ""), z.get("jedinica", "")]
            wc.writerow([store_id, pid, z.get("cijena", ""), z.get("jedinicna", ""), z.get("posebna", ""),
                         z.get("akcija", ""), izvor, z["_redak"]])
    p = io.StringIO()
    wp = csv.writer(p)
    wp.writerow(["product_id", "barcode", "name", "brand", "quantity", "unit"])
    wp.writerows(proizvodi.values())
    return trgovine.getvalue(), p.getvalue(), cijene.getvalue(), izvori


def main():
    d = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else datetime.now(ZoneInfo("Europe/Zagreb")).date()
    odabrani = sys.argv[2:] or list(LANCI)
    for lanac in odabrani:
        t = time.time()
        rez = preuzmi_lanac(lanac, d)
        redaka = sum(len(x["zapisi"]) for x in rez.get("koristene", []))
        greske = [x for x in rez["datoteke"] if x["greska"]]
        print(json.dumps({"lanac": lanac, "uspjeh": rez["uspjeh"], "razlog": rez["razlog"],
                          "datoteka": len(rez["datoteke"]), "poslovnica": len(rez.get("koristene", [])), "redaka": redaka,
                          "neuspjelih_datoteka": len(greske), "primjer_greske": greske[0]["greska"] if greske else "",
                          "nepoznati_stupci": sorted({n for x in rez.get("koristene", []) for n in x["nepoznati"]})[:12],
                          "sekundi": round(time.time() - t)}, ensure_ascii=False), flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
